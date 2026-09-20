
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import os
import re
from typing import Dict, Tuple, Optional

import pandas as pd
import requests

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


REQUIRED_COLUMNS = [
    "Week Ending",
    "Store",
    "Region",
    "Workshops",
    "Brand Presentations",
    "Revenue",
]

NAVY = RGBColor(0, 72, 56)
TEAL = RGBColor(0, 126, 72)
LIGHT = RGBColor(231, 242, 234)
DARK = RGBColor(17, 17, 17)
MUTED = RGBColor(64, 87, 79)
WHITE = RGBColor(255, 255, 255)
WARM = RGBColor(247, 244, 238)
MID_GREY = RGBColor(185, 198, 192)
PROJECT_DIR = Path(__file__).resolve().parent
MANAGEMENT_PROMPT_PATH = PROJECT_DIR / "prompts" / "management_summary_prompt.txt"


def load_management_prompt() -> str:
    """Load the governed AI instructions used for management commentary."""
    return MANAGEMENT_PROMPT_PATH.read_text(encoding="utf-8").strip()


def _extract_date_from_filename(name: str) -> Optional[pd.Timestamp]:
    match = re.search(r"(\d{4}-\d{2}-\d{2})", name)
    if not match:
        return None
    try:
        return pd.Timestamp(match.group(1)).normalize()
    except Exception:
        return None


def validate_week_file(path: Path) -> Dict:
    """Validate one weekly Excel file and return metadata for the UI."""
    result = {
        "file": path.name,
        "valid": False,
        "week_ending": None,
        "rows": 0,
        "message": "",
    }

    try:
        df = pd.read_excel(path, sheet_name=0)
    except Exception as exc:
        result["message"] = f"Could not read file: {exc}"
        return result

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        result["message"] = f"Missing columns: {', '.join(missing)}"
        return result

    df = df.copy()
    df["Week Ending"] = pd.to_datetime(df["Week Ending"], errors="coerce").dt.normalize()

    if df["Week Ending"].isna().any():
        result["message"] = "Week Ending contains an invalid date."
        return result

    weeks = df["Week Ending"].dropna().unique()
    if len(weeks) != 1:
        result["message"] = "Each weekly file must contain exactly one Week Ending date."
        return result

    for col in ["Workshops", "Brand Presentations", "Revenue"]:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.isna().any():
            result["message"] = f"{col} contains a non-numeric value."
            return result
        if (converted < 0).any():
            result["message"] = f"{col} contains a negative value."
            return result

    if df["Store"].astype(str).str.strip().eq("").any():
        result["message"] = "Store contains a blank value."
        return result

    if df.duplicated(subset=["Week Ending", "Store"]).any():
        result["message"] = "Duplicate store rows were found for the same week."
        return result

    week_ending = pd.Timestamp(weeks[0]).normalize()
    filename_date = _extract_date_from_filename(path.name)
    if filename_date is not None and filename_date != week_ending:
        result["message"] = (
            f"Filename date {filename_date.date()} does not match "
            f"Week Ending {week_ending.date()}."
        )
        return result

    result.update({
        "valid": True,
        "week_ending": week_ending,
        "rows": len(df),
        "message": "Validated",
    })
    return result


def load_history(data_dir: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Load and append all validated weekly workbooks in the data folder."""
    data_dir = Path(data_dir)
    files = sorted(data_dir.glob("*.xlsx"))
    if not files:
        raise FileNotFoundError(f"No Excel files found in {data_dir}")

    frames = []
    statuses = []

    for path in files:
        status = validate_week_file(path)
        statuses.append(status)
        if not status["valid"]:
            continue

        df = pd.read_excel(path, sheet_name=0)
        df["Week Ending"] = pd.to_datetime(df["Week Ending"]).dt.normalize()
        for col in ["Workshops", "Brand Presentations", "Revenue"]:
            df[col] = pd.to_numeric(df[col])
        df["Source File"] = path.name
        frames.append(df)

    if not frames:
        raise ValueError("No valid weekly files are available.")

    history = pd.concat(frames, ignore_index=True)
    history = history.sort_values(["Week Ending", "Store"]).reset_index(drop=True)

    # Defensive check across files.
    if history.duplicated(subset=["Week Ending", "Store"]).any():
        raise ValueError(
            "Duplicate Week Ending + Store records exist across the validated files."
        )

    status_df = pd.DataFrame(statuses)
    return history, status_df


def calculate_kpis(history: pd.DataFrame, selected_week) -> Dict:
    selected_week = pd.Timestamp(selected_week).normalize()
    available = sorted(pd.to_datetime(history["Week Ending"]).dt.normalize().unique())
    available = [pd.Timestamp(x).normalize() for x in available]

    if selected_week not in available:
        raise ValueError("Selected week is not available in the historical data.")

    previous_candidates = [w for w in available if w < selected_week]
    previous_week = max(previous_candidates) if previous_candidates else None

    current = history[history["Week Ending"] == selected_week].copy()
    previous = (
        history[history["Week Ending"] == previous_week].copy()
        if previous_week is not None
        else pd.DataFrame(columns=history.columns)
    )

    def snapshot(df: pd.DataFrame) -> Dict:
        if df.empty:
            return {
                "workshops": 0,
                "brand_presentations": 0,
                "total_activities": 0,
                "active_stores": 0,
                "revenue": 0.0,
                "revenue_per_store": 0.0,
            }
        activity = df["Workshops"] + df["Brand Presentations"]
        active_stores = int(df.loc[activity > 0, "Store"].nunique())
        revenue = float(df["Revenue"].sum())
        return {
            "workshops": int(df["Workshops"].sum()),
            "brand_presentations": int(df["Brand Presentations"].sum()),
            "total_activities": int(activity.sum()),
            "active_stores": active_stores,
            "revenue": revenue,
            "revenue_per_store": revenue / active_stores if active_stores else 0.0,
        }

    cur = snapshot(current)
    prev = snapshot(previous)

    def pct_change(a, b):
        if b in (0, 0.0, None):
            return None
        return (a - b) / b

    store_activity = (
        current.assign(Total_Activity=current["Workshops"] + current["Brand Presentations"])
        .groupby("Store", as_index=False)
        .agg(
            Workshops=("Workshops", "sum"),
            Brand_Presentations=("Brand Presentations", "sum"),
            Revenue=("Revenue", "sum"),
            Total_Activity=("Total_Activity", "sum"),
        )
        .sort_values(["Total_Activity", "Revenue"], ascending=False)
    )

    top_store = store_activity.iloc[0]["Store"] if not store_activity.empty else "N/A"
    top_store_activity = (
        int(store_activity.iloc[0]["Total_Activity"]) if not store_activity.empty else 0
    )
    region_revenue = (
        current.groupby("Region", as_index=False)["Revenue"]
        .sum()
        .sort_values("Revenue", ascending=False)
    )
    top_region = region_revenue.iloc[0]["Region"] if not region_revenue.empty else "N/A"
    top_region_revenue = (
        float(region_revenue.iloc[0]["Revenue"]) if not region_revenue.empty else 0.0
    )
    weekly_trend = (
        history[history["Week Ending"] <= selected_week]
        .groupby("Week Ending", as_index=False)
        .agg(
            Workshops=("Workshops", "sum"),
            Brand_Presentations=("Brand Presentations", "sum"),
            Revenue=("Revenue", "sum"),
        )
        .sort_values("Week Ending")
        .tail(8)
        .reset_index(drop=True)
    )
    weekly_trend["Total_Activities"] = (
        weekly_trend["Workshops"] + weekly_trend["Brand_Presentations"]
    )

    return {
        "selected_week": selected_week,
        "previous_week": previous_week,
        "current": cur,
        "previous": prev,
        "changes": {
            "workshops": cur["workshops"] - prev["workshops"] if previous_week is not None else None,
            "brand_presentations": (
                cur["brand_presentations"] - prev["brand_presentations"]
                if previous_week is not None else None
            ),
            "total_activities": (
                cur["total_activities"] - prev["total_activities"]
                if previous_week is not None else None
            ),
            "active_stores": (
                cur["active_stores"] - prev["active_stores"]
                if previous_week is not None else None
            ),
            "revenue": cur["revenue"] - prev["revenue"] if previous_week is not None else None,
            "revenue_pct": pct_change(cur["revenue"], prev["revenue"]) if previous_week is not None else None,
            "revenue_per_store_pct": (
                pct_change(cur["revenue_per_store"], prev["revenue_per_store"])
                if previous_week is not None else None
            ),
            "workshops_pct": pct_change(cur["workshops"], prev["workshops"]) if previous_week is not None else None,
            "brand_pct": (
                pct_change(cur["brand_presentations"], prev["brand_presentations"])
                if previous_week is not None else None
            ),
        },
        "top_store": top_store,
        "top_store_activity": top_store_activity,
        "top_region": top_region,
        "top_region_revenue": top_region_revenue,
        "weekly_trend": weekly_trend,
        "store_activity": store_activity,
        "current_df": current,
        "previous_df": previous,
    }


def deterministic_summary(kpis: Dict) -> str:
    """Fallback summary used when no OpenRouter key is configured."""
    c = kpis["current"]
    ch = kpis["changes"]
    prev_week = kpis["previous_week"]

    lines = []
    if prev_week is None:
        lines.append(
            f"• Revenue for the reporting period was ${c['revenue']:,.0f}."
        )
        lines.append(
            f"• This is the first available reporting period, with {c['total_activities']} total additional activities."
        )
        lines.append(f"• Team education sessions totalled {c['workshops']}.")
        lines.append(f"• In-store brand activations totalled {c['brand_presentations']}.")
    else:
        def direction(value):
            if value > 0:
                return "increased"
            if value < 0:
                return "decreased"
            return "was unchanged"

        if ch["revenue_pct"] is not None:
            lines.append(
                f"• Revenue {direction(ch['revenue'])} by "
                f"{abs(ch['revenue_pct']) * 100:.1f}% to ${c['revenue']:,.0f}."
            )
        lines.append(
            f"• Total additional activities {direction(ch['total_activities'])} from "
            f"{kpis['previous']['total_activities']} to {c['total_activities']}."
        )
        lines.append(
            f"• Team education sessions {direction(ch['workshops'])} from "
            f"{kpis['previous']['workshops']} to {c['workshops']}."
        )
        lines.append(
            f"• In-store brand activations {direction(ch['brand_presentations'])} from "
            f"{kpis['previous']['brand_presentations']} to {c['brand_presentations']}."
        )

    lines.append(
        f"• {kpis['top_store']} recorded the highest activity with "
        f"{kpis['top_store_activity']} education sessions and brand activations combined."
    )
    return "\n".join(lines[:5])


def _build_ai_prompt(kpis: Dict) -> str:
    c = kpis["current"]
    p = kpis["previous"]
    ch = kpis["changes"]
    top = kpis["store_activity"].head(5)

    top_rows = "\n".join(
        f"- {r.Store}: team_education_sessions={int(r.Workshops)}, "
        f"in_store_brand_activations={int(r.Brand_Presentations)}, "
        f"revenue=${r.Revenue:,.0f}"
        for r in top.itertuples(index=False)
    )

    previous_label = (
        kpis["previous_week"].strftime("%d %b %Y")
        if kpis["previous_week"] is not None
        else "not available"
    )

    return f"""
Produce the weekly management summary using the governed system instructions and
the validated reporting data below.

Reporting week: {kpis['selected_week'].strftime('%d %b %Y')}
Previous week: {previous_label}

Current week:
- Team education sessions: {c['workshops']}
- In-store brand activations: {c['brand_presentations']}
- Total additional activities: {c['total_activities']}
- Active stores: {c['active_stores']}
- Revenue: ${c['revenue']:,.0f}

Previous week:
- Team education sessions: {p['workshops']}
- In-store brand activations: {p['brand_presentations']}
- Total additional activities: {p['total_activities']}
- Active stores: {p['active_stores']}
- Revenue: ${p['revenue']:,.0f}

Validated changes:
- Team education sessions change: {ch['workshops']}
- In-store brand activations change: {ch['brand_presentations']}
- Total additional activities change: {ch['total_activities']}
- Active stores change: {ch['active_stores']}
- Revenue change: {ch['revenue']}
- Revenue percent change: {None if ch['revenue_pct'] is None else round(ch['revenue_pct'] * 100, 1)}%

Top stores this week:
{top_rows}
""".strip()


def generate_management_summary(
    kpis: Dict,
    api_key: Optional[str] = None,
    model: str = "openai/gpt-4o-mini",
) -> Tuple[str, str]:
    """
    Return (summary, source).
    source is 'OpenRouter' or 'Built-in fallback'.
    """
    api_key = api_key or os.getenv("OPENROUTER_API_KEY")
    if not api_key:
        return deterministic_summary(kpis), "Built-in fallback"

    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": load_management_prompt(),
            },
            {"role": "user", "content": _build_ai_prompt(kpis)},
        ],
        "temperature": 0.2,
        # Reasoning models need enough room for both internal reasoning and the
        # visible answer. Keep reasoning light so the summary remains concise.
        "max_completion_tokens": 1200,
        "reasoning": {"effort": "low", "exclude": True},
    }

    try:
        response = requests.post(
            "https://openrouter.ai/api/v1/chat/completions",
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=30,
        )
        response.raise_for_status()
        result = response.json()
        choice = result["choices"][0]
        content = (choice.get("message") or {}).get("content")
        if not isinstance(content, str) or not content.strip():
            finish_reason = choice.get("finish_reason", "unknown")
            raise ValueError(
                f"Model returned no visible text (finish_reason={finish_reason})"
            )
        return content.strip(), "OpenRouter"
    except Exception:
        # The demo must still work during an interview even if the model/API is unavailable.
        return deterministic_summary(kpis), "Built-in fallback"


def _clean_summary_lines(summary: str):
    lines = []
    for raw in summary.splitlines():
        line = raw.strip()
        if not line:
            continue
        line = re.sub(r"^[\-\*\u2022]+\s*", "", line)
        line = re.sub(r"^\d+[\.\)]\s*", "", line)
        if line:
            lines.append(line)
    return lines[:6]


def _add_title(slide, title: str, subtitle: str = ""):
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(12.0), Inches(0.55))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.67), Inches(0.98), Inches(11.6), Inches(0.35))
        p2 = sub.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED


def _add_footer(slide, text="Weekly Reporting Automation Demo"):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.0),
        Inches(7.18),
        Inches(13.333),
        Inches(0.32),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = NAVY
    line.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.55), Inches(7.22), Inches(12.2), Inches(0.18))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(8)
    p.font.color.rgb = WHITE


def _add_kpi_card(slide, x, y, w, h, label, value, delta=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(222, 228, 233)

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.14)

    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = MUTED

    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(25)
    p2.font.bold = True
    p2.font.color.rgb = NAVY

    if delta:
        p3 = tf.add_paragraph()
        p3.text = delta
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEAL


def _delta_text(value, unit=""):
    if value is None:
        return "No previous week"
    sign = "+" if value > 0 else ""
    return f"{sign}{value}{unit} vs previous week"


def _add_ai_callout(slide, text: str, y=6.15):
    """Add one short, governed AI insight beneath a chart."""
    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.72),
        Inches(y),
        Inches(11.9),
        Inches(0.68),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = LIGHT
    band.line.color.rgb = RGBColor(201, 224, 209)
    tf = band.text_frame
    tf.clear()
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.10)
    p = tf.paragraphs[0]
    lead = p.add_run()
    lead.text = "AI commentary  "
    lead.font.size = Pt(11)
    lead.font.bold = True
    lead.font.color.rgb = TEAL
    body = p.add_run()
    body.text = text
    body.font.size = Pt(11)
    body.font.color.rgb = DARK


def _style_chart(chart, colors, data_label_position=XL_LABEL_POSITION.OUTSIDE_END):
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.color.rgb = MUTED
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.color.rgb = DARK
    for series, color in zip(chart.series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.format.line.color.rgb = color
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = data_label_position
    labels.show_value = True
    labels.font.size = Pt(9)
    labels.font.color.rgb = DARK


def generate_powerpoint(kpis: Dict, summary: str) -> bytes:
    """Create a chart-led executive PPTX with governed AI commentary."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    week_label = kpis["selected_week"].strftime("%d %B %Y")
    prev_label = (
        kpis["previous_week"].strftime("%d %B %Y")
        if kpis["previous_week"] is not None
        else "Not available"
    )
    c, ch = kpis["current"], kpis["changes"]
    lines = _clean_summary_lines(summary)
    if not lines:
        lines = _clean_summary_lines(deterministic_summary(kpis))
    while len(lines) < 5:
        lines.append("")

    # Slide 1: minimal cover
    slide = prs.slides.add_slide(blank)
    block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    block.fill.solid()
    block.fill.fore_color.rgb = NAVY
    block.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.15), Inches(0.12), Inches(3.25)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(1.15), Inches(1.25), Inches(10.8), Inches(2.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Weekly performance review"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = f"Week ending {week_label}"
    p2.font.size = Pt(19)
    p2.font.color.rgb = RGBColor(215, 229, 237)
    p3 = tf.add_paragraph()
    p3.text = "Power BI and validated Excel reporting"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(185, 205, 216)

    # Slide 2: executive overview
    slide = prs.slides.add_slide(blank)
    _add_title(
        slide,
        "Executive overview",
        f"Week ending {week_label} compared with {prev_label}",
    )
    rev_delta = (
        "No prior comparison"
        if ch["revenue_pct"] is None
        else f"{ch['revenue_pct'] * 100:+.1f}% from previous week"
    )
    _add_kpi_card(
        slide, Inches(0.72), Inches(1.55), Inches(3.0), Inches(1.65),
        "Revenue", f"${c['revenue']/1000:,.0f}K", rev_delta
    )
    _add_kpi_card(
        slide, Inches(3.94), Inches(1.55), Inches(2.7), Inches(1.65),
        "Total additional activities", f"{c['total_activities']}", _delta_text(ch["total_activities"])
    )
    _add_kpi_card(
        slide, Inches(6.86), Inches(1.55), Inches(2.7), Inches(1.65),
        "Active stores", f"{c['active_stores']}", _delta_text(ch["active_stores"])
    )
    _add_kpi_card(
        slide, Inches(9.78), Inches(1.55), Inches(2.82), Inches(1.65),
        "Leading activity store", kpis["top_store"], f"{kpis['top_store_activity']} activities"
    )
    section = slide.shapes.add_textbox(Inches(0.78), Inches(3.70), Inches(11.7), Inches(1.65))
    tf = section.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "What changed this week"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    for text in lines[:2]:
        pp = tf.add_paragraph()
        pp.text = text
        pp.font.size = Pt(16)
        pp.font.color.rgb = DARK
        pp.space_before = Pt(10)
    _add_ai_callout(slide, lines[0], y=6.02)
    _add_footer(slide)

    # Slide 3: revenue trend
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Revenue trend", "Validated Power BI measure across available reporting weeks")
    trend = kpis["weekly_trend"]
    categories = [pd.Timestamp(v).strftime("%d %b") for v in trend["Week Ending"]]

    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series("Revenue ($000)", [round(v / 1000, 1) for v in trend["Revenue"]])

    graphic = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.82),
        Inches(1.55),
        Inches(11.7),
        Inches(4.25),
        chart_data,
    )
    chart = graphic.chart
    _style_chart(chart, [TEAL])
    chart.has_legend = False
    chart.value_axis.tick_labels.number_format = '$0"K"'
    chart.plots[0].data_labels.number_format = '$0"K"'
    _add_ai_callout(slide, lines[0])
    _add_footer(slide)

    # Slide 4: activity trend
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Additional activity trend", "Validated weekly Excel measures after integration into the reporting model")
    activity_data = CategoryChartData()
    activity_data.categories = categories
    activity_data.add_series("Team education sessions", [int(v) for v in trend["Workshops"]])
    activity_data.add_series("In-store brand activations", [int(v) for v in trend["Brand_Presentations"]])
    graphic = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED,
        Inches(0.82), Inches(1.55), Inches(11.7), Inches(4.25), activity_data,
    )
    chart = graphic.chart
    _style_chart(chart, [TEAL, RGBColor(140, 196, 158)], XL_LABEL_POSITION.INSIDE_END)
    _add_ai_callout(slide, f"{lines[1]} {lines[2]}")
    _add_footer(slide)

    # Slide 5: top-store comparison
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Store activity comparison", f"Top five stores for week ending {week_label}")
    top = kpis["store_activity"].head(5).iloc[::-1]
    store_data = CategoryChartData()
    store_data.categories = [str(v) for v in top["Store"]]
    store_data.add_series("Team education sessions", [int(v) for v in top["Workshops"]])
    store_data.add_series("In-store brand activations", [int(v) for v in top["Brand_Presentations"]])
    graphic = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_STACKED,
        Inches(0.82), Inches(1.55), Inches(11.7), Inches(4.25), store_data,
    )
    chart = graphic.chart
    _style_chart(chart, [TEAL, RGBColor(140, 196, 158)], XL_LABEL_POSITION.CENTER)
    _add_ai_callout(slide, lines[4])
    _add_footer(slide)

    # Slide 6: AI-assisted management summary
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Management summary", "AI-assisted wording grounded in validated KPIs")

    summary_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.78),
        Inches(1.55),
        Inches(11.75),
        Inches(4.8),
    )
    summary_box.fill.solid()
    summary_box.fill.fore_color.rgb = LIGHT
    summary_box.line.color.rgb = RGBColor(222, 228, 233)
    tf = summary_box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.35)
    tf.margin_right = Inches(0.35)
    tf.margin_top = Inches(0.30)

    for i, line in enumerate(lines[:5]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.alignment = PP_ALIGN.LEFT
        p.font.size = Pt(17)
        p.font.color.rgb = DARK
        p.space_after = Pt(16)
    note = slide.shapes.add_textbox(Inches(0.95), Inches(6.48), Inches(11.2), Inches(0.35))
    p = note.text_frame.paragraphs[0]
    p.text = "Control: KPI calculations remain outside the language model. Commentary requires analyst review."
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    _add_footer(slide)

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out.getvalue()
